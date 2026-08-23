# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

JP = "Meiryo"
YEN  = '¥#,##0;▲¥#,##0;"-"'
YEN0 = '¥#,##0;▲¥#,##0;"¥0"'
YENFREE = '¥#,##0;▲¥#,##0;"¥0　（無償提供）"'
QTY = '#,##0;▲#,##0;"-"'

BLUE  = Font(name=JP, size=10, color="0000FF")           # 手入力
GREEN = Font(name=JP, size=10, color="008000")           # 他シート参照
BLACK = Font(name=JP, size=10)
BOLD  = Font(name=JP, size=10, bold=True)
YFILL = PatternFill("solid", fgColor="FFFF00")
HFILL = PatternFill("solid", fgColor="D9D9D9")
LFILL = PatternFill("solid", fgColor="F2F2F2")
AFILL = PatternFill("solid", fgColor="FFF2CC")

thin  = Side(style="thin",  color="808080")
med   = Side(style="medium", color="000000")
BOX      = Border(left=thin, right=thin, top=thin, bottom=thin)
BOX_MED  = Border(left=med,  right=med,  top=med,  bottom=med)

def merge(ws, rng, value=None, font=None, align=None, fill=None, border=None, fmt=None):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    if value is not None: c.value = value
    if font: c.font = font
    if align: c.alignment = align
    if fmt: c.number_format = fmt
    first, last = rng.split(":")
    c0, r0 = ws[first].column, ws[first].row
    c1, r1 = ws[last].column,  ws[last].row
    for r in range(r0, r1 + 1):
        for cc in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=cc)
            if fill: cell.fill = fill
            if border: cell.border = border
    return c

L  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
Lt = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
C  = Alignment(horizontal="center", vertical="center", wrap_text=True)
R  = Alignment(horizontal="right",  vertical="center")

wb = Workbook()

# ============================================================
# シート1：基本設定
# ============================================================
ws = wb.active
ws.title = "基本設定"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [2.5, 26, 52, 4, 60]):
    ws.column_dimensions[col].width = w

ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = "B1:E36"
merge(ws, "B2:C2", "基本設定（最初にここを入力してください）", Font(name=JP, size=14, bold=True), L)
merge(ws, "B3:C3", "黄色のセルを入力すると、すべての見積書シートに自動で反映されます。", Font(name=JP, size=9, color="808080"), L)

rows = [
    (5,  "■ 発行者情報", None, None),
    (6,  "会社名",           "【会社名】", "例：株式会社〇〇"),
    (7,  "住所",             "〒000-0000 【住所】", "例：〒150-0001 東京都渋谷区〇〇1-2-3"),
    (8,  "電話番号",         "000-0000-0000", ""),
    (9,  "メールアドレス",   "【メールアドレス】", ""),
    (10, "担当者名",         "【担当者名】", ""),
    (11, "振込先",           "【銀行名・支店名・種別・口座番号・口座名義】", "追加ご発注が発生した場合のみ使用します"),
    (13, "■ 宛先・見積情報", None, None),
    (14, "宛先（お客様名）", "〇〇〇〇", "相手ごとに書き換えてください"),
    (15, "敬称",             "御中", "法人＝御中／個人＝様"),
    (16, "発行日",           "2026年8月23日", ""),
    (17, "見積番号",         "MG-2026-0001", "相手ごとに連番で振ってください"),
    (18, "見積有効期限",     "発行日より30日間", ""),
    (19, "納期",             "ご発注後 約4〜6週間（ヒアリング完了より起算）", ""),
    (21, "■ 計算条件", None, None),
    (22, "消費税率",         0.1, "税率が変わった場合はここだけ変更"),
    (23, "通常価格（4話＋総集編・税別）", 250000, "無料提供の定価。備考の記載にも使われます"),
]
for r, label, val, note in rows:
    if val is None:
        ws.cell(row=r, column=2, value=label).font = Font(name=JP, size=11, bold=True)
        continue
    lc = ws.cell(row=r, column=2, value=label)
    lc.font = BOLD; lc.fill = LFILL; lc.alignment = L; lc.border = BOX
    vc = ws.cell(row=r, column=3, value=val)
    vc.font = BLUE; vc.fill = YFILL; vc.alignment = L; vc.border = BOX
    if label == "消費税率": vc.number_format = "0.0%"
    if "通常価格" in label: vc.number_format = YEN
    if note:
        ws.cell(row=r, column=5, value="← " + note).font = Font(name=JP, size=9, color="808080")

merge(ws, "B25:C25", "凡例", Font(name=JP, size=11, bold=True), L)
lg = [
    ("黄色のセル", "入力欄。ここだけ書き換えてください。"),
    ("青文字",     "手入力の値。"),
    ("緑文字",     "他シートを参照している値。直接書き換えないでください。"),
    ("黒文字",     "自動計算されるセル。触らないでください。"),
]
for i, (k, v) in enumerate(lg):
    r = 26 + i
    a = ws.cell(row=r, column=2, value=k); a.font = BOLD; a.fill = LFILL; a.alignment = L; a.border = BOX
    b = ws.cell(row=r, column=3, value=v); b.font = BLACK; b.alignment = L; b.border = BOX
ws["B26"].font = Font(name=JP, size=10, bold=True); ws["B26"].fill = YFILL
ws["B27"].font = Font(name=JP, size=10, bold=True, color="0000FF")
ws["B28"].font = Font(name=JP, size=10, bold=True, color="008000")

merge(ws, "B31:C31", "使い方", Font(name=JP, size=11, bold=True), L)
steps = [
    "① このシートの黄色いセルを埋める（会社情報は一度だけ。宛先・見積番号は相手ごとに変更）",
    "② 「見積書（無料制作）」シートを印刷／PDF出力して相手に送付する",
    "③ 追加のご依頼が発生したら「見積書（追加オプション）」シートの数量欄に数を入れて送付する",
    "④ 単価を変えたいときは「料金表」シートを編集する（追加オプション見積書に自動反映）",
]
for i, s in enumerate(steps):
    merge(ws, f"B{32+i}:C{32+i}", s, BLACK, L)

# ============================================================
# 共通：見積書レイアウト
# ============================================================
def quote_header(ws, subject_text, payterm_text):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [2.5, 8, 41, 8, 8, 14, 16]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.5; ws.page_margins.right = 0.5
    ws.page_margins.top = 0.5;  ws.page_margins.bottom = 0.5

    ws.row_dimensions[2].height = 34
    merge(ws, "B2:G2", "御 見 積 書", Font(name=JP, size=22, bold=True), C)

    # 宛先
    ws.row_dimensions[4].height = 24
    a = merge(ws, "B4:D4", "='基本設定'!$C$14&\"　\"&'基本設定'!$C$15",
              Font(name=JP, size=14, bold=True, color="008000"), L)
    for cc in range(2, 5):
        ws.cell(row=4, column=cc).border = Border(bottom=med)

    ws.cell(row=4, column=6, value="発行日").font = BOLD
    ws.cell(row=4, column=6).alignment = R
    v = ws.cell(row=4, column=7, value="='基本設定'!$C$16"); v.font = GREEN; v.alignment = R
    ws.cell(row=5, column=6, value="見積番号").font = BOLD
    ws.cell(row=5, column=6).alignment = R
    v = ws.cell(row=5, column=7, value="='基本設定'!$C$17"); v.font = GREEN; v.alignment = R

    merge(ws, "B6:D6", "下記の通りお見積り申し上げます。", BLACK, L)

    # 発行者ブロック
    merge(ws, "E7:G7",  "='基本設定'!$C$6",  Font(name=JP, size=12, bold=True, color="008000"), L)
    merge(ws, "E8:G8",  "='基本設定'!$C$7",  GREEN, L)
    merge(ws, "E9:G9",  "=\"TEL：\"&'基本設定'!$C$8",  GREEN, L)
    merge(ws, "E10:G10","=\"Email：\"&'基本設定'!$C$9", GREEN, L)
    merge(ws, "E11:G11","=\"担当：\"&'基本設定'!$C$10", GREEN, L)
    for r in range(7, 12):
        for cc in range(5, 8):
            ws.cell(row=r, column=cc).border = Border(
                left=thin if cc == 5 else None, right=thin if cc == 7 else None,
                top=thin if r == 7 else None,   bottom=thin if r == 11 else None)

    # 左ブロック
    left = [(8, "件　名", subject_text, False),
            (9, "有効期限", "='基本設定'!$C$18", True),
            (10, "納　期", "='基本設定'!$C$19", True),
            (11, "お支払条件", payterm_text, False)]
    for r, label, val, is_ref in left:
        lc = ws.cell(row=r, column=2, value=label)
        lc.font = Font(name=JP, size=9, bold=True); lc.fill = LFILL
        lc.alignment = Alignment(horizontal="center", vertical="center"); lc.border = BOX
        vc = merge(ws, f"C{r}:D{r}", val, GREEN if is_ref else BLACK, L, border=BOX)

def table_header(ws, row):
    ws.row_dimensions[row].height = 22
    heads = [("B", "No."), ("C", "品目・仕様"), ("D", "数量"), ("E", "単位"), ("F", "単価"), ("G", "金額")]
    for col, h in heads:
        c = ws[f"{col}{row}"]
        c.value = h; c.font = BOLD; c.fill = HFILL; c.alignment = C; c.border = BOX

def totals(ws, first_item_row, last_item_row, start_row):
    labels = [("小　計（税別）", f"=SUM(G{first_item_row}:G{last_item_row})", False),
              ("=\"消費税（\"&TEXT('基本設定'!$C$22,\"0%\")&\"）\"", f"=ROUND(G{start_row}*'基本設定'!$C$22,0)", False),
              ("合計金額（税込）", f"=G{start_row}+G{start_row+1}", True)]
    for i, (label, formula, is_total) in enumerate(labels):
        r = start_row + i
        ws.row_dimensions[r].height = 22 if not is_total else 26
        lc = merge(ws, f"C{r}:F{r}", label,
                   Font(name=JP, size=12, bold=True) if is_total else BOLD, R,
                   fill=AFILL if is_total else LFILL, border=BOX)
        ws.cell(row=r, column=2).border = BOX
        ws.cell(row=r, column=2).fill = AFILL if is_total else LFILL
        vc = ws.cell(row=r, column=7, value=formula)
        vc.font = Font(name=JP, size=12, bold=True) if is_total else BOLD
        vc.alignment = R; vc.number_format = YEN0
        vc.border = BOX_MED if is_total else BOX
        if is_total: vc.fill = AFILL
    return start_row + 2

def amount_box(ws, row, total_row, fmt=None):
    ws.row_dimensions[row].height = 34
    merge(ws, f"B{row}:D{row}", "御 見 積 金 額（税込）",
          Font(name=JP, size=13, bold=True), C, fill=AFILL, border=BOX_MED)
    merge(ws, f"E{row}:G{row}", f"=G{total_row}",
          Font(name=JP, size=18, bold=True), R, border=BOX_MED, fmt=fmt or YEN0)

def notes(ws, row, lines, title="備　考"):
    merge(ws, f"B{row}:G{row}", title, BOLD, L, fill=HFILL, border=BOX)
    for i, t in enumerate(lines):
        r = row + 1 + i
        ws.row_dimensions[r].height = max(15, 12.5 * -(-len(t) // 46))
        merge(ws, f"B{r}:G{r}", t, Font(name=JP, size=9), Lt, border=BOX)
    return row + len(lines)

# ============================================================
# シート2：見積書（無料制作）
# ============================================================
ws = wb.create_sheet("見積書（無料制作）")
quote_header(ws, "社長ショートドラマ制作　全4話＋総集編　一式",
             "本見積は無償提供のため、お支払いはございません")

amount_box(ws, 13, 26, YENFREE)
table_header(ws, 15)

items = [
    ("企画・構成（全4話＋総集編／エピソード設計・構成台本）", 1, "式", 40000),
    ("脚本制作（本編 全4話／ナレーション・セリフ・テロップ原稿）", 4, "話", 15000),
    ("映像制作・編集（本編 全4話／AI映像生成・カット編集）", 4, "話", 25000),
    ("総集編 制作・編集（1本）", 1, "本", 25000),
    ("テロップ挿入・BGM・効果音・仕上げ", 1, "式", 15000),
    ("各SNS向け書き出し・納品（Instagram / TikTok / YouTube Shorts）", 1, "式", 5000),
    ("修正対応（軽微修正 1回分）", 1, "回", 5000),
]
r = 16
for i, (name, qty, unit, price) in enumerate(items, start=1):
    ws.row_dimensions[r].height = max(20, 15 * -(-len(name) // 27))
    ws.cell(row=r, column=2, value=i).font = BLACK
    ws.cell(row=r, column=2).alignment = C
    ws.cell(row=r, column=3, value=name).font = BLACK
    ws.cell(row=r, column=3).alignment = L
    ws.cell(row=r, column=4, value=qty).font = BLACK
    ws.cell(row=r, column=4).alignment = C
    ws.cell(row=r, column=4).number_format = QTY
    ws.cell(row=r, column=5, value=unit).font = BLACK
    ws.cell(row=r, column=5).alignment = C
    ws.cell(row=r, column=6, value=price).font = BLACK
    ws.cell(row=r, column=6).alignment = R
    ws.cell(row=r, column=6).number_format = YEN
    ws.cell(row=r, column=7, value=f"=D{r}*F{r}").font = BLACK
    ws.cell(row=r, column=7).alignment = R
    ws.cell(row=r, column=7).number_format = YEN
    for cc in range(2, 8):
        ws.cell(row=r, column=cc).border = BOX
    r += 1

# 割引行
ws.row_dimensions[r].height = 30
ws.cell(row=r, column=2, value=len(items) + 1).font = BOLD
ws.cell(row=r, column=2).alignment = C
ws.cell(row=r, column=3, value="特別ご協力割引（本企画へのご参加条件のご履行を条件とする全額割引）").font = Font(name=JP, size=10, bold=True, color="FF0000")
ws.cell(row=r, column=3).alignment = L
ws.cell(row=r, column=4, value=1).font = BLACK
ws.cell(row=r, column=4).alignment = C
ws.cell(row=r, column=5, value="式").font = BLACK
ws.cell(row=r, column=5).alignment = C
ws.cell(row=r, column=6, value=f"=-SUM(G16:G{r-1})").font = Font(name=JP, size=10, bold=True, color="FF0000")
ws.cell(row=r, column=6).alignment = R
ws.cell(row=r, column=6).number_format = YEN
ws.cell(row=r, column=7, value=f"=D{r}*F{r}").font = Font(name=JP, size=10, bold=True, color="FF0000")
ws.cell(row=r, column=7).alignment = R
ws.cell(row=r, column=7).number_format = YEN
for cc in range(2, 8):
    ws.cell(row=r, column=cc).border = BOX

disc_row = r
total_end = totals(ws, 16, disc_row, disc_row + 1)

note_lines = [
    "※本見積書は、下記4条件のご履行を前提とした特別無償提供です。",
    "　①納品後7日以内に、ご本人アカウントにて投稿（Instagram「共同投稿者」機能で弊社アカウントを共同投稿者に設定／6ヶ月間の掲載継続）",
    "　②キャプションへの制作クレジット記載および弊社公式アカウントのタグ付け・メンション",
    "　③第1話公開時・総集編公開時の計2回、ストーリーズでのご紹介（弊社アカウントのメンション付き）",
    "　④完成動画の弊社二次利用（実績紹介・公式アカウントでの再投稿・営業資料・広告素材）のご許可",
    "=\"※上記条件が履行されない場合、通常価格 \"&TEXT('基本設定'!$C$23,\"¥#,##0\")&\"（税別）を申し受ける場合がございます。\"",
    "※無償修正は納品後1回・軽微修正のみとし、納品後3日以内に一括でご依頼ください。範囲を超える修正は「料金表」に基づき別途お見積りいたします。",
    "　【軽微修正に含まれるもの】テロップの誤字・表記ゆれ／事実誤認の訂正／BGM・効果音の音量調整／尺の微調整（±3秒程度）／ご提供済み素材の差し替え",
    "　【別途お見積りとなるもの】構成・脚本の変更／シーン・カットの作り直し／登場人物のビジュアル変更／演出方向の変更／全体の作り直し",
    "※脚本のご確認をもって内容の確定といたします。映像制作着手後の構成変更は別途お見積りとなります。",
    "※確認のご依頼から3営業日以内にご返信がない場合は、ご承認いただいたものとみなし次工程へ進行いたします。",
    "※素材（お写真・ロゴ等）のご提出はご依頼から7日以内にお願いいたします。2週間ご連絡が取れない場合はいったん制作終了とし、再開は別途お見積りといたします。",
    "※本ドラマの映像はAIにより生成しております。登場人物は実在の人物ではありません。ご本人のお写真を使用する場合は事前にご相談いたします。",
    "※エピソード内の実績・数値等、事実関係のご確認はご本人にてお願いいたします。",
    "※制作途中でご中止となった場合、制作済みデータの権利は弊社に帰属し、公開はご遠慮いただきます。",
]
notes(ws, total_end + 2, note_lines)

# ============================================================
# シート3：見積書（追加オプション）
# ============================================================
ws = wb.create_sheet("見積書（追加オプション）")
quote_header(ws, "社長ショートドラマ制作　追加オプション",
             "納品後 翌月末までに指定口座へお振込み（振込手数料は貴社ご負担）")

amount_box(ws, 13, 25)
table_header(ws, 15)

# 料金表を参照する6行
for i in range(6):
    r = 16 + i
    src = 4 + i
    ws.row_dimensions[r].height = 20
    ws.cell(row=r, column=2, value=i + 1).font = BLACK
    ws.cell(row=r, column=2).alignment = C
    ws.cell(row=r, column=3, value=f"='料金表'!$C${src}").font = GREEN
    ws.cell(row=r, column=3).alignment = L
    q = ws.cell(row=r, column=4, value=0)
    q.font = BLUE; q.fill = YFILL; q.alignment = C; q.number_format = QTY
    ws.cell(row=r, column=5, value=f"='料金表'!$D${src}").font = GREEN
    ws.cell(row=r, column=5).alignment = C
    ws.cell(row=r, column=6, value=f"='料金表'!$E${src}").font = GREEN
    ws.cell(row=r, column=6).alignment = R
    ws.cell(row=r, column=6).number_format = YEN
    ws.cell(row=r, column=7, value=f"=D{r}*F{r}").font = BLACK
    ws.cell(row=r, column=7).alignment = R
    ws.cell(row=r, column=7).number_format = YEN
    for cc in range(2, 8):
        ws.cell(row=r, column=cc).border = BOX

# 特急対応割増
r = 22
ws.row_dimensions[r].height = 20
ws.cell(row=r, column=2, value=7).font = BLACK
ws.cell(row=r, column=2).alignment = C
ws.row_dimensions[r].height = 30
ws.cell(row=r, column=3, value="特急対応割増（ご依頼から72時間以内の納品／上記小計の50%）").font = BLACK
ws.cell(row=r, column=3).alignment = L
q = ws.cell(row=r, column=4, value=0)
q.font = BLUE; q.fill = YFILL; q.alignment = C; q.number_format = QTY
ws.cell(row=r, column=5, value="式").font = BLACK
ws.cell(row=r, column=5).alignment = C
ws.cell(row=r, column=6, value="=ROUND(SUM(G16:G21)*0.5,0)").font = BLACK
ws.cell(row=r, column=6).alignment = R
ws.cell(row=r, column=6).number_format = YEN
ws.cell(row=r, column=7, value=f"=D{r}*F{r}").font = BLACK
ws.cell(row=r, column=7).alignment = R
ws.cell(row=r, column=7).number_format = YEN
for cc in range(2, 8):
    ws.cell(row=r, column=cc).border = BOX

opt_end = totals(ws, 16, 22, 23)

opt_notes = [
    "※黄色のセル（数量欄）にご希望の数量をご入力ください。金額・小計・消費税・合計は自動で計算されます。",
    "※単価は「料金表」シートと連動しています。単価を変更する場合は「料金表」シートをご編集ください。",
    "※特急対応割増は、数量欄に「1」を入力すると項番1〜6の小計の50%が加算されます。",
    "",
    "【記入例】追加修正を2回、シーン再生成を3カットご依頼の場合",
    "　　項番5「追加修正（軽微修正・2回目以降）」の数量欄に「2」→ 金額 ¥20,000",
    "　　項番4「シーン再生成」の数量欄に「3」→ 金額 ¥15,000",
    "　　小計 ¥35,000 ／ 消費税 ¥3,500 ／ 合計 ¥38,500 が自動表示されます",
    "",
    "※無償制作分（全4話＋総集編）に含まれる範囲を超えるご依頼が対象です。",
    "※作業着手はご承認のご返信をいただいた後となります。",
    "※お支払いは納品後、翌月末までに指定口座へお振込みをお願いいたします。",
    "※本見積の有効期限は発行日より30日間です。",
]
notes(ws, opt_end + 2, opt_notes, "備　考・ご記入方法")


# ============================================================
# シート4：料金表
# ============================================================
ws = wb.create_sheet("料金表")
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = "B1:F17"
for col, w in zip("ABCDEF", [2.5, 5, 48, 8, 14, 46]):
    ws.column_dimensions[col].width = w

ws.row_dimensions[1].height = 26
merge(ws, "B1:F1", "追加オプション 料金表（税別）", Font(name=JP, size=16, bold=True), L)
merge(ws, "B2:F2", "黄色の単価セルを変更すると「見積書（追加オプション）」シートに自動反映されます。",
      Font(name=JP, size=9, color="808080"), L)

ws.row_dimensions[3].height = 22
for col, h in [("B", "No."), ("C", "項目"), ("D", "単位"), ("E", "単価（税別）"), ("F", "備考")]:
    c = ws[f"{col}3"]
    c.value = h; c.font = BOLD; c.fill = HFILL; c.alignment = C; c.border = BOX

rate_rows = [
    ("社長ショートドラマ 4話＋総集編パッケージ", "式", 250000, "企画・脚本・映像制作・編集・納品まで一式（有料でお受けする場合の通常価格）"),
    ("追加話数の制作", "話", 50000, "第5話以降。企画・脚本・映像制作込み"),
    ("構成・脚本からの作り直し", "話", 30000, "映像制作着手後の構成変更・脚本の書き直し"),
    ("シーン再生成", "カット", 5000, "1カットあたり。既存構成のまま映像のみ差し替え"),
    ("追加修正（軽微修正・2回目以降）", "回", 10000, "納品後1回目の軽微修正は無償。2回目以降が対象"),
    ("横型・正方形など別比率の書き出し", "本", 5000, "1本あたり。16:9／1:1 などへのリサイズ書き出し"),
]
for i, (name, unit, price, note) in enumerate(rate_rows):
    r = 4 + i
    ws.row_dimensions[r].height = 20
    ws.cell(row=r, column=2, value=i + 1).font = BLACK
    ws.cell(row=r, column=2).alignment = C
    ws.cell(row=r, column=3, value=name).font = BLACK
    ws.cell(row=r, column=3).alignment = L
    ws.cell(row=r, column=4, value=unit).font = BLACK
    ws.cell(row=r, column=4).alignment = C
    p = ws.cell(row=r, column=5, value=price)
    p.font = BLUE; p.fill = YFILL; p.alignment = R; p.number_format = YEN
    ws.cell(row=r, column=6, value=note).font = Font(name=JP, size=9)
    ws.cell(row=r, column=6).alignment = L
    for cc in range(2, 7):
        ws.cell(row=r, column=cc).border = BOX

r = 10
ws.row_dimensions[r].height = 20
ws.cell(row=r, column=2, value=7).font = BLACK
ws.cell(row=r, column=2).alignment = C
ws.cell(row=r, column=3, value="特急対応割増（ご依頼から72時間以内の納品）").font = BLACK
ws.cell(row=r, column=3).alignment = L
ws.cell(row=r, column=4, value="式").font = BLACK
ws.cell(row=r, column=4).alignment = C
ws.cell(row=r, column=5, value="対象小計の50%").font = BLACK
ws.cell(row=r, column=5).alignment = R
ws.cell(row=r, column=6, value="見積書シート側で自動計算されます（単価の手入力は不要）").font = Font(name=JP, size=9)
ws.cell(row=r, column=6).alignment = L
for cc in range(2, 7):
    ws.cell(row=r, column=cc).border = BOX

merge(ws, "B12:F12", "料金設定の考え方（社内メモ・相手には渡さないでください）", Font(name=JP, size=11, bold=True), L)
memo = [
    "・通常価格 250,000円は「4話＋総集編」を有料で受注する場合の価格。無料提供の価値を相手に示すための基準値でもある。",
    "・追加修正 10,000円/回は、金額そのものより「無限に依頼できない」ことを示すための設定。安すぎると抑止にならず、高すぎると角が立つ水準。",
    "・構成・脚本からの作り直し 30,000円/話は、映像制作着手後の手戻りコストに相当。ここを無料にすると際限なく巻き戻される。",
    "・断るのではなく「有料でお受けできます」と選ばせる形にすることで、細かい注文が売上に変わる。値引きはしない。",
    "・相手によって単価を変えない。スムーズに進めてくれた方が損をする運用にすると、紹介が止まる。",
]
for i, t in enumerate(memo):
    merge(ws, f"B{13+i}:F{13+i}", t, Font(name=JP, size=9), Lt)

# ============================================================
wb.active = 1
out = "/home/user/My-labo2-_/docs/社長ショートドラマ_御見積書.xlsx"
wb.save(out)
print("saved:", out)
